# -*- coding: utf-8 -*-
"""




implement(STT), imagedescription(Vision), contentdescriptionand content.
corresponds to tasks.md P2.1.
"""

from __future__ import annotations

import asyncio
import base64
import mimetypes
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Any, Optional, Protocol, Union

from pydantic import BaseModel, Field


class MediaType(Enum):
    """media type"""
    AUDIO = "audio"
    IMAGE = "image"
    VIDEO = "video"
    DOCUMENT = "document"
    UNKNOWN = "unknown"


class STTProvider(Enum):
    """provider"""
    OPENAI = "openai"
    GROQ = "groq"
    GOOGLE = "google"
    MISTRAL = "mistral"
    DEEPGRAM = "deepgram"
    MINIMAX = "minimax"


class VisionProvider(Enum):
    """provider"""
    OPENAI = "openai"
    GOOGLE = "google"
    ANTHROPIC = "anthropic"


@dataclass
class MediaContent:
    """

content
    
    Attributes:
        media_type:media type
        path:file path
        mime_type:MIME type
        data:count(optional)
        url:URL(optional)
        size_bytes:
        duration_seconds:()
        metadata:additional metadata
    
"""
    media_type: MediaType
    path: str = ""
    mime_type: str = ""
    data: Optional[bytes] = None
    url: Optional[str] = None
    size_bytes: int = 0
    duration_seconds: float = 0.0
    metadata: dict[str, Any] = field(default_factory=dict)
    
    @classmethod
    def from_path(cls, path: str) -> "MediaContent":
        """fromfile pathcreate"""
        p = Path(path)
        mime_type, _ = mimetypes.guess_type(path)
        mime_type = mime_type or "application/octet-stream"
        
        # media type
        media_type = MediaType.UNKNOWN
        if mime_type.startswith("audio/"):
            media_type = MediaType.AUDIO
        elif mime_type.startswith("image/"):
            media_type = MediaType.IMAGE
        elif mime_type.startswith("video/"):
            media_type = MediaType.VIDEO
        elif mime_type.startswith("application/pdf") or mime_type.startswith("text/"):
            media_type = MediaType.DOCUMENT
        
        size_bytes = p.stat().st_size if p.exists() else 0
        
        return cls(
            media_type=media_type,
            path=path,
            mime_type=mime_type,
            size_bytes=size_bytes,
        )
    
    def to_base64(self) -> str:
        """base64 characters"""
        if self.data:
            return base64.b64encode(self.data).decode("utf-8")
        if self.path and Path(self.path).exists():
            with open(self.path, "rb") as f:
                return base64.b64encode(f.read()).decode("utf-8")
        return ""


@dataclass
class UnderstandingResult:
    """Understanding result"""
    success: bool
    text: str = ""
    confidence: float = 1.0
    provider: str = ""
    language: str = ""
    duration_ms: int = 0
    metadata: dict[str, Any] = field(default_factory=dict)
    error: Optional[str] = None


class MediaUnderstandingProvider(ABC):
    """provider"""
    
    @abstractmethod
    async def understand(self, content: MediaContent) -> UnderstandingResult:
        """Understand media content"""
        ...
    
    @property
    @abstractmethod
    def supported_types(self) -> list[MediaType]:
        """supportmedia type"""
        ...


class OpenAISTTProvider(MediaUnderstandingProvider):
    """Speech-to-text provider backed by OpenAI Whisper."""
    
    def __init__(
        self,
        api_key: str,
        model: str = "whisper-1",
        language: Optional[str] = None,
    ) -> None:
        self.api_key = api_key
        self.model = model
        self.language = language
    
    @property
    def supported_types(self) -> list[MediaType]:
        return [MediaType.AUDIO]
    
    async def understand(self, content: MediaContent) -> UnderstandingResult:
        """Transcribe audio content with the OpenAI Whisper API."""
        import time
        import httpx
        
        start = time.monotonic()
        
        try:
            # Resolve the source audio bytes.
            if content.data:
                file_data = content.data
            elif content.path and Path(content.path).exists():
                with open(content.path, "rb") as f:
                    file_data = f.read()
            else:
                return UnderstandingResult(
                    success=False,
                    error="No audio data provided",
                    provider="openai",
                )
            
            # Call the OpenAI transcription API.
            async with httpx.AsyncClient() as client:
                files = {"file": ("audio.mp3", file_data, content.mime_type)}
                data = {"model": self.model}
                if self.language:
                    data["language"] = self.language
                
                response = await client.post(
                    "https://api.openai.com/v1/audio/transcriptions",
                    headers={"Authorization": f"Bearer {self.api_key}"},
                    files=files,
                    data=data,
                    timeout=120.0,
                )
                
                if response.status_code != 200:
                    return UnderstandingResult(
                        success=False,
                        error=f"API error: {response.status_code} - {response.text}",
                        provider="openai",
                    )
                
                result = response.json()
                text = result.get("text", "")
                
                return UnderstandingResult(
                    success=True,
                    text=text,
                    provider="openai",
                    language=self.language or "auto",
                    duration_ms=int((time.monotonic() - start) * 1000),
                )
                
        except Exception as e:
            return UnderstandingResult(
                success=False,
                error=str(e),
                provider="openai",
                duration_ms=int((time.monotonic() - start) * 1000),
            )


class OpenAIVisionProvider(MediaUnderstandingProvider):
    """Image understanding provider backed by OpenAI vision models."""
    
    def __init__(
        self,
        api_key: str,
        model: str = "gpt-4o",
        prompt: str = "Please describe the content of this image.",
        max_tokens: int = 500,
    ) -> None:
        self.api_key = api_key
        self.model = model
        self.prompt = prompt
        self.max_tokens = max_tokens
    
    @property
    def supported_types(self) -> list[MediaType]:
        return [MediaType.IMAGE]
    
    async def understand(self, content: MediaContent) -> UnderstandingResult:
        """Describe image content with an OpenAI vision model."""
        import time
        import httpx
        
        start = time.monotonic()
        
        try:
            # Build the image payload from a URL or inline base64 data.
            if content.url:
                image_content = {"type": "image_url", "image_url": {"url": content.url}}
            else:
                base64_data = content.to_base64()
                if not base64_data:
                    return UnderstandingResult(
                        success=False,
                        error="No image data provided",
                        provider="openai",
                    )
                image_content = {
                    "type": "image_url",
                    "image_url": {"url": f"data:{content.mime_type};base64,{base64_data}"}
                }
            
            # Call the OpenAI chat completions API for vision analysis.
            async with httpx.AsyncClient() as client:
                response = await client.post(
                    "https://api.openai.com/v1/chat/completions",
                    headers={
                        "Authorization": f"Bearer {self.api_key}",
                        "Content-Type": "application/json",
                    },
                    json={
                        "model": self.model,
                        "messages": [
                            {
                                "role": "user",
                                "content": [
                                    {"type": "text", "text": self.prompt},
                                    image_content,
                                ],
                            }
                        ],
                        "max_tokens": self.max_tokens,
                    },
                    timeout=60.0,
                )
                
                if response.status_code != 200:
                    return UnderstandingResult(
                        success=False,
                        error=f"API error: {response.status_code} - {response.text}",
                        provider="openai",
                    )
                
                result = response.json()
                text = result["choices"][0]["message"]["content"]
                
                return UnderstandingResult(
                    success=True,
                    text=text,
                    provider="openai",
                    duration_ms=int((time.monotonic() - start) * 1000),
                )
                
        except Exception as e:
            return UnderstandingResult(
                success=False,
                error=str(e),
                provider="openai",
                duration_ms=int((time.monotonic() - start) * 1000),
            )


class DocumentExtractor(MediaUnderstandingProvider):
    """content"""
    
    def __init__(self, max_chars: int = 50000) -> None:
        self.max_chars = max_chars
    
    @property
    def supported_types(self) -> list[MediaType]:
        return [MediaType.DOCUMENT]
    
    async def understand(self, content: MediaContent) -> UnderstandingResult:
        """content"""
        import time
        
        start = time.monotonic()
        
        try:
            path = Path(content.path)
            if not path.exists():
                return UnderstandingResult(
                    success=False,
                    error=f"File not found: {content.path}",
                    provider="document_extractor",
                )
            
            # based on type
            suffix = path.suffix.lower()
            text = ""
            
            if suffix == ".pdf":
                text = await self._extract_pdf(path)
            elif suffix in (".txt", ".md", ".json", ".csv", ".xml", ".html"):
                text = await self._extract_text(path)
            elif suffix in (".docx", ".doc"):
                text = await self._extract_docx(path)
            elif suffix in (".xlsx", ".xls"):
                text = await self._extract_excel(path)
            else:
                return UnderstandingResult(
                    success=False,
                    error=f"Unsupported document type: {suffix}",
                    provider="document_extractor",
                )
            
            # truncate
            if len(text) > self.max_chars:
                text = text[:self.max_chars] + "\n...[Content truncated]"
            
            return UnderstandingResult(
                success=True,
                text=text,
                provider="document_extractor",
                duration_ms=int((time.monotonic() - start) * 1000),
                metadata={"truncated": len(text) >= self.max_chars},
            )
            
        except Exception as e:
            return UnderstandingResult(
                success=False,
                error=str(e),
                provider="document_extractor",
                duration_ms=int((time.monotonic() - start) * 1000),
            )
    
    async def _extract_text(self, path: Path) -> str:
        """"""
        return path.read_text(encoding="utf-8", errors="ignore")
    
    async def _extract_pdf(self, path: Path) -> str:
        """PDF content"""
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            text_parts = []
            for page in reader.pages:
                text_parts.append(page.extract_text() or "")
            return "\n".join(text_parts)
        except ImportError:
            return f"[PDF extraction requires pypdf library: {path.name}]"
    
    async def _extract_docx(self, path: Path) -> str:
        """Word content"""
        try:
            import docx
            doc = docx.Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            return f"[Word extraction requires python-docx library: {path.name}]"
    
    async def _extract_excel(self, path: Path) -> str:
        """Excel content"""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True)
            text_parts = []
            for sheet in wb.worksheets:
                text_parts.append(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(c) if c is not None else "" for c in row)
                    text_parts.append(row_text)
            return "\n".join(text_parts)
        except ImportError:
            return f"[Excel 提取需要 openpyxl 库: {path.name}]"


class MediaUnderstandingHandler:
    """

handle
    
    .
    
    Example usage::
    
        handler = MediaUnderstandingHandler()
        handler.register_provider(OpenAISTTProvider(api_key="..."))
        handler.register_provider(OpenAIVisionProvider(api_key="..."))
        
        content = MediaContent.from_path("audio.mp3")
        result = await handler.understand(content)
        print(result.text)
    
"""
    
    def __init__(self) -> None:
        self._providers: dict[MediaType, list[MediaUnderstandingProvider]] = {}
    
    def register_provider(self, provider: MediaUnderstandingProvider) -> None:
        """registerprovider"""
        for media_type in provider.supported_types:
            if media_type not in self._providers:
                self._providers[media_type] = []
            self._providers[media_type].append(provider)
    
    async def understand(
        self,
        content: MediaContent,
        *,
        fallback: bool = True,
    ) -> UnderstandingResult:
        """

Understand media content
        
        Args:
            content:content
            fallback:at provider
            
        Returns:
            Understanding result
        
"""
        providers = self._providers.get(content.media_type, [])
        
        if not providers:
            return UnderstandingResult(
                success=False,
                error=f"No provider for media type: {content.media_type.value}",
            )
        
        last_error = None
        for provider in providers:
            result = await provider.understand(content)
            if result.success:
                return result
            last_error = result.error
            if not fallback:
                break
        
        return UnderstandingResult(
            success=False,
            error=last_error or "All providers failed",
        )
    
    def supports(self, media_type: MediaType) -> bool:
        """check supportmedia type"""
        return media_type in self._providers and len(self._providers[media_type]) > 0


def create_media_handler(
    *,
    openai_api_key: Optional[str] = None,
    enable_document_extraction: bool = True,
) -> MediaUnderstandingHandler:
    """

create handle factory count
    
    Args:
        openai_api_key:OpenAI API key
        enable_document_extraction:
        
    Returns:
        configuration handle
    
"""
    handler = MediaUnderstandingHandler()
    
    if openai_api_key:
        handler.register_provider(OpenAISTTProvider(api_key=openai_api_key))
        handler.register_provider(OpenAIVisionProvider(api_key=openai_api_key))
    
    if enable_document_extraction:
        handler.register_provider(DocumentExtractor())
    
    return handler
